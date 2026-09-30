"""Excel break report for an operations team: what is open, what is new, what cleared, and the evidence for each."""

import argparse
import sqlite3
import sys
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src import answers, config, db, holdings, ingest, staging

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
ZAR = '"R"#,##0.00;("R"#,##0.00);-'
PCT = "0.00%;(0.00%);-"
INT = "#,##0"
DATE = "yyyy-mm-dd"
DEC = "0.0"

# Bucket labels include "days" because Excel reads a bare "2-5" in a COUNTIFS criterion as the 5th of February
AGE_BUCKET = '=IF({c}{r}<=1,"0-1 days",IF({c}{r}<=5,"2-5 days",IF({c}{r}<=20,"6-20 days","21+ days")))'
BUCKETS = ["0-1 days", "2-5 days", "6-20 days", "21+ days"]

BREAK_CODES = [
    ("MATCH", "Both sides agree within tolerance"),
    ("VAL", "Both sides have a price and they differ by more than both the rand floor and the percentage"),
    ("ONE_A", "Only side A has a price for this security and date"),
    ("ONE_B", "Only side B has a price; in a restatement recon, a price published late"),
    ("DUP", "A side has more than one row for the same security and date"),
    ("UNIT", "The sides are about 100x apart, or one side's unit is flagged or unknown"),
    ("CAL", "A price on a day the JSE was closed"),
    ("STALE", "A value break where one side's price has not moved for 5 or more sessions"),
    ("TIMING", "A break that cleared within 2 trading days of first being seen"),
]


def as_date(text):
    return date.fromisoformat(text) if text else None


def as_of(conn: sqlite3.Connection) -> str:
    return conn.execute("SELECT MAX(snapshot_date) FROM ingest_run").fetchone()[0]


def open_breaks(conn):
    return conn.execute(
        """
        SELECT e.recon_name, e.key_id, m.name, m.sector, e.price_date, e.first_seen, e.last_seen, e.age_days,
               e.latest_status, x.close_a, x.close_b, x.diff_zar, x.diff_pct / 100.0, e.latest_explanation,
               n.resolution, n.note
        FROM break_episode e
        LEFT JOIN security_master m ON m.security_id = e.key_id
        LEFT JOIN recon_run r ON r.recon_name = e.recon_name AND r.snapshot_b = e.last_seen
        LEFT JOIN recon_result x
            ON x.recon_run_id = r.recon_run_id AND x.key_id = e.key_id AND x.price_date = e.price_date
        LEFT JOIN break_note n ON n.recon_name = e.recon_name AND n.key_id = e.key_id AND n.price_date = e.price_date
        WHERE e.state = 'OPEN'
        ORDER BY e.age_days DESC, e.recon_name, e.key_id, e.price_date
        """
    ).fetchall()


def new_breaks(conn):
    return conn.execute(
        """
        SELECT e.recon_name, e.key_id, m.name, e.price_date, e.latest_status, x.close_a, x.close_b,
               x.diff_zar, x.diff_pct / 100.0, e.latest_explanation, e.state
        FROM break_episode e
        LEFT JOIN security_master m ON m.security_id = e.key_id
        LEFT JOIN recon_run r ON r.recon_name = e.recon_name AND r.snapshot_b = e.first_seen
        LEFT JOIN recon_result x
            ON x.recon_run_id = r.recon_run_id AND x.key_id = e.key_id AND x.price_date = e.price_date
        WHERE e.first_seen = (SELECT MAX(snapshot_b) FROM recon_run WHERE recon_name = e.recon_name)
        ORDER BY e.recon_name, e.key_id, e.price_date
        """
    ).fetchall()


def cleared_breaks(conn):
    return conn.execute(
        """
        SELECT e.recon_name, e.key_id, m.name, e.price_date, e.first_seen, e.cleared_on, e.observations,
               e.first_status, e.age_days, e.state, e.latest_explanation, n.resolution, n.note
        FROM break_episode e
        LEFT JOIN security_master m ON m.security_id = e.key_id
        LEFT JOIN break_note n ON n.recon_name = e.recon_name AND n.key_id = e.key_id AND n.price_date = e.price_date
        WHERE e.state <> 'OPEN'
        ORDER BY e.cleared_on DESC, e.recon_name, e.key_id
        """
    ).fetchall()


def restatements(conn):
    run = conn.execute(
        "SELECT recon_run_id, snapshot_a, snapshot_b FROM recon_run WHERE recon_name LIKE '%\\_restatement' ESCAPE '\\' "
        "ORDER BY snapshot_b DESC LIMIT 1"
    ).fetchone()
    if not run:
        return None, []
    rows = conn.execute(
        "SELECT key_id, price_date, status, close_a, close_b, diff_zar, explanation FROM recon_result "
        "WHERE recon_run_id = ? AND status <> 'MATCH' ORDER BY price_date, key_id",
        (run[0],),
    ).fetchall()
    return run, rows


def dq_rows(conn):
    rows = []
    for source, snap, _ in staging.latest_snapshots(conn):
        if config.sources().get(source, {}).get("daily_batch"):
            continue  # a rotated source's snapshot covers only a few securities
        f = staging.findings(conn, source, snap)
        for day, sec, kind, close, ref1, ref2 in f["anomalies"]:
            side = "below" if kind == "too_small" else "above"
            rows.append(("Unit anomaly", sec, day, f"R{close:,.2f}, about 100x {side} its nearest bars "
                         f"(R{ref1:,.2f} and R{ref2:,.2f})", source))
        for day, expected, missing in f["whole"]:
            rows.append(("Missing session", "All", day, f"{missing} of {expected} securities have no price", source))
        for day, sec in f["missing_bars"]:
            rows.append(("Missing price", sec, day, "No price on a trading day", source))
        for sec, end, days, close in f["stale"]:
            rows.append(("Frozen price", sec, end, f"R{close:,.2f} unchanged for {days} sessions", source))
        for day, reason, bars, traded in f["closed"]:
            rows.append(("Price on closed day", "All", day, f"{bars} prices on {reason}, {traded} with volume", source))
    return rows


def holding_rows(conn):
    latest = holdings.latest_file(conn)
    if not latest:
        return None, []
    rows = conn.execute(
        """
        SELECT key_id, status, broker_lines, broker_value, broker_price, implied_qty, book_qty, settled_qty,
               our_price, prev_price, book_value, value_diff, explanation
        FROM holding_recon_result WHERE file_id = ?
        ORDER BY status = 'MATCH', ABS(COALESCE(value_diff, 1e12)) DESC, key_id
        """,
        (latest[0],),
    ).fetchall()
    return latest, rows


def style(ws):
    for row in ws.iter_rows():
        for cell in row:
            bold = cell.font.bold if cell.font else False
            color = cell.font.color if cell.font else None
            italic = cell.font.italic if cell.font else False
            cell.font = Font(name=FONT, size=10, bold=bold, italic=italic, color=color)


def table(ws, columns, rows, empty_note, formulas=None):
    ws.append([c[0] for c in columns])
    for cell in ws[1]:
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r, values in enumerate(rows, start=2):
        values = [as_date(v) if fmt == DATE else v for v, (_, _, fmt) in zip(values, columns)]
        ws.append(values)
        for col, template in (formulas or {}).items():
            ws[f"{col}{r}"] = template.format(r=r)
    if not rows:
        ws["A2"] = empty_note
        ws["A2"].font = Font(name=FONT, italic=True, color="808080")
    for i, (_, width, fmt) in enumerate(columns, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = width
        if fmt:
            for r in range(2, len(rows) + 2):
                ws[f"{letter}{r}"].number_format = fmt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(len(rows) + 1, 2)}"
    ws.row_dimensions[1].height = 30
    style(ws)


def summary(ws, day, generated):
    ws["A1"] = "JSE price reconciliation: break report"
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    ws["A2"] = f"As of {day}. Generated {generated}. Every count below is a formula over the tab named beside it."
    lines = [
        ("Breaks between sources", None, None),
        ("Open breaks", "=SUM(B6:B9)", "Open Breaks"),
        *[(f"  open {b}", f"=COUNTIFS('Open Breaks'!$I:$I,\"{b}\")", "Open Breaks") for b in BUCKETS],
        ("New in the latest comparison", "=COUNTA('New Today'!$B:$B)-1", "New Today"),
        ("Cleared as timing differences", '=COUNTIFS(Cleared!$J:$J,"TIMING")', "Cleared"),
        ("Cleared after longer", '=COUNTIFS(Cleared!$J:$J,"CLEARED")', "Cleared"),
        ("Mean trading days to clear", '=IFERROR(AVERAGE(Cleared!$I:$I),"-")', "Cleared"),
        ("Restatements since the previous snapshot", None, None),
        ("Prices rewritten", '=COUNTIFS(Restatements!$C:$C,"VAL")+COUNTIFS(Restatements!$C:$C,"UNIT")'
                             '+COUNTIFS(Restatements!$C:$C,"STALE")', "Restatements"),
        ("Prices published late", '=COUNTIFS(Restatements!$C:$C,"ONE_B")', "Restatements"),
        ("Prices withdrawn", '=COUNTIFS(Restatements!$C:$C,"ONE_A")', "Restatements"),
        ("Data quality, latest full snapshot", None, None),
        ("Unit anomalies", '=COUNTIFS(\'Data Quality\'!$A:$A,"Unit anomaly")', "Data Quality"),
        ("Missing sessions", '=COUNTIFS(\'Data Quality\'!$A:$A,"Missing session")', "Data Quality"),
        ("Missing prices", '=COUNTIFS(\'Data Quality\'!$A:$A,"Missing price")', "Data Quality"),
        ("Frozen prices", '=COUNTIFS(\'Data Quality\'!$A:$A,"Frozen price")', "Data Quality"),
        ("Prices on closed days", '=COUNTIFS(\'Data Quality\'!$A:$A,"Price on closed day")', "Data Quality"),
        ("Broker statement against the internal book", None, None),
        ("Positions compared", "=COUNTA(Holdings!$B:$B)-1", "Holdings"),
        ("Positions agreeing", '=COUNTIFS(Holdings!$B:$B,"MATCH")', "Holdings"),
        ("Holdings breaks", '=COUNTA(Holdings!$B:$B)-1-COUNTIFS(Holdings!$B:$B,"MATCH")', "Holdings"),
        ("Net difference (R)", "=SUM(Holdings!$L:$L)", "Holdings"),
        ("Gross difference (R)", "=SUMPRODUCT(ABS(Holdings!$L$2:$L$2000))", "Holdings"),
    ]
    for r, (label, formula, tab) in enumerate(lines, start=4):
        ws[f"A{r}"] = label
        if formula is None:
            for col in "ABC":
                ws[f"{col}{r}"].fill = SECTION_FILL
            ws[f"A{r}"].font = Font(name=FONT, bold=True)
            continue
        ws[f"B{r}"] = formula
        ws[f"B{r}"].number_format = DEC if "AVERAGE" in formula else ZAR if "(R)" in label else INT
        ws[f"C{r}"] = f"{tab} tab"
    assert ws["A5"].value == "Open breaks" and ws["A6"].value.startswith("  open 0-1")  # B5 sums B6:B9
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 18
    style(ws)
    ws["A1"].font = Font(name=FONT, size=14, bold=True)


def about(ws, conn, day, generated):
    tol = config.tolerance_rules()
    rows = [("As of", day), ("Generated", generated)]
    for source, snap, cutoff in staging.latest_snapshots(conn):
        rows.append((f"Latest {source} snapshot", f"{snap}, sessions complete to {cutoff}"))
    for name, rules in tol["price_recon"].items():
        c = rules["close"]
        rows.append((f"Tolerance, {name}", f"A break needs more than R{c['abs_floor_zar']:.2f} and "
                                          f"{c['rel_pct']:.2f}% difference in the close"))
    rows += [
        ("Age", "Trading days on the JSE calendar, from first sighting to clearing or to the latest comparison"),
        ("TIMING", f"Cleared within {tol['dq']['timing_clear_days']} trading days of first being seen"),
        ("Rotation", "EODHD compares about 18 securities a day, so a break can only clear when its security "
                     "comes round again, roughly weekly"),
        ("EODHD unit", "Assumed cents (ZAc); the API reports none. ABG.JSE closed 21743.0 on 2026-09-28, the "
                       "same as Yahoo's 21743 ZAc"),
        ("Materiality", "ASISA Standard on NAV calculation for CIS portfolios, s10.3.3: suggested maximum "
                        "tolerance for a pricing error is 0.5% of NAV. https://asisa.org.za/media/om5kbuwu/"
                        "asisa-standard-nav-calculation-for-cis-portfolios-november-2015.pdf"),
        ("NAV basis", "Costs in the Answers tab assume an equal-weighted fund across the securities priced "
                      "that day; illustrative, not a real holding"),
        ("Notes", "Resolutions and analyst notes come from config/break_notes.yaml in the repository"),
        ("", ""),
        ("Break codes", ""),
        *BREAK_CODES,
    ]
    for r, (key, value) in enumerate(rows, start=1):
        ws[f"A{r}"], ws[f"B{r}"] = key, value
        ws[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 110
    style(ws)
    for r, (key, _) in enumerate(rows, start=1):
        if key == "Break codes":
            ws[f"A{r}"].font = Font(name=FONT, bold=True)


def build(conn: sqlite3.Connection, path) -> None:
    day = as_of(conn)
    generated = datetime.now(ingest.SAST).strftime("%Y-%m-%d %H:%M SAST")
    none = f"None as of {day}"
    wb = Workbook()
    summary(wb.active, day, generated)
    wb.active.title = "Summary"

    table(wb.create_sheet("Open Breaks"), [
        ("Recon", 16, None), ("Security", 10, None), ("Name", 26, None), ("Sector", 20, None),
        ("Price date", 12, DATE), ("First seen", 12, DATE), ("Last seen", 12, DATE),
        ("Age (trading days)", 11, INT), ("Age bucket", 11, None), ("Status", 9, None),
        ("Close A (R)", 13, ZAR), ("Close B (R)", 13, ZAR), ("Difference (R)", 13, ZAR),
        ("Difference (%)", 12, PCT), ("Explanation", 60, None), ("Resolution", 16, None), ("Note", 40, None),
    ], [(*r[:8], None, *r[8:]) for r in open_breaks(conn)], none, formulas={"I": AGE_BUCKET.format(c="H", r="{r}")})

    table(wb.create_sheet("New Today"), [
        ("Recon", 16, None), ("Security", 10, None), ("Name", 26, None), ("Price date", 12, DATE),
        ("Status", 9, None), ("Close A (R)", 13, ZAR), ("Close B (R)", 13, ZAR), ("Difference (R)", 13, ZAR),
        ("Difference (%)", 12, PCT), ("Explanation", 60, None), ("State", 10, None),
    ], new_breaks(conn), none)

    table(wb.create_sheet("Cleared"), [
        ("Recon", 16, None), ("Security", 10, None), ("Name", 26, None), ("Price date", 12, DATE),
        ("First seen", 12, DATE), ("Cleared on", 12, DATE), ("Comparisons", 12, INT), ("First status", 11, None),
        ("Age (trading days)", 11, INT), ("State", 10, None), ("Last explanation", 60, None),
        ("Resolution", 16, None), ("Note", 40, None),
    ], cleared_breaks(conn), none)

    run, rows = restatements(conn)
    table(wb.create_sheet("Restatements"), [
        ("Security", 10, None), ("Price date", 12, DATE), ("Status", 9, None),
        (f"Close in {run[1]} snapshot (R)" if run else "Earlier close (R)", 16, ZAR),
        (f"Close in {run[2]} snapshot (R)" if run else "Later close (R)", 16, ZAR),
        ("Difference (R)", 13, ZAR), ("Explanation", 60, None),
    ], rows, none if run else "No restatement comparison yet: it needs two snapshots of the same source")

    table(wb.create_sheet("Data Quality"), [
        ("Check", 20, None), ("Security", 10, None), ("Date", 12, DATE), ("Detail", 70, None), ("Source", 10, None),
    ], dq_rows(conn), none)

    latest, rows = holding_rows(conn)
    table(wb.create_sheet("Holdings"), [
        ("Security", 10, None), ("Status", 9, None), ("Statement lines", 10, INT), ("Broker value (R)", 15, ZAR),
        ("Broker price (R)", 13, ZAR), ("Implied quantity", 13, "#,##0.0000"), ("Book quantity", 13, "#,##0.0000"),
        ("Settled quantity", 13, "#,##0.0000"), ("Our close (R)", 13, ZAR), ("Previous close (R)", 13, ZAR),
        ("Book at our close (R)", 15, ZAR), ("Difference (R)", 14, ZAR), ("Explanation", 70, None),
    ], rows, "No broker statement loaded")
    if latest:
        wb["Holdings"]["O1"] = ("SYNTHETIC statement" if latest[2] else "Statement") + f" as of {latest[1]}"
        wb["Holdings"]["O1"].font = Font(name=FONT, bold=True, color="C00000")

    ws = wb.create_sheet("Answers")
    ws.column_dimensions["A"].width = 150
    for source in answers.feeds_under_test(conn):
        for line in answers.answers(conn, source):
            ws.append([line])
            ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    style(ws)
    ws["A1"].font = Font(name=FONT, size=12, bold=True)

    about(wb.create_sheet("About"), conn, day, generated)
    wb.save(path)


def main(argv=None) -> int:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--out", help="output path; defaults to reports/break_report_<as of>.xlsx")
    args = p.parse_args(argv)
    conn = db.connect()
    db.apply_schema(conn)
    path = args.out or config.ROOT / "reports" / f"break_report_{as_of(conn)}.xlsx"
    build(conn, path)
    print(f"Wrote {path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
