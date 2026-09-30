"""The Excel break report: tabs, rows, and formulas that point at the columns they claim to count."""

import re
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from src import config, db, lifecycle, report, trading_calendar
from tests.test_lifecycle import D1, D2, D3, D6, DAY, NAME

SHEETS = ["Summary", "Open Breaks", "New Today", "Cleared", "Restatements", "Data Quality", "Holdings", "Approved Prices",
          "Answers", "About"]


class ReportTest(unittest.TestCase):
    def setUp(self):
        tmp = Path(tempfile.mkdtemp())
        self.path = tmp / "report.xlsx"
        self.conn = db.connect(tmp / "test.db")
        self.addCleanup(self.conn.close)
        db.apply_schema(self.conn)
        trading_calendar.load(self.conn)
        self.conn.execute("INSERT INTO security_master (security_id, name, sector) VALUES ('OPN', 'Open Ltd', 'Test')")
        self.conn.execute(
            """
            INSERT INTO ingest_run (run_id, source, snapshot_date, fetched_at, session_cutoff, lookback_period,
                symbols_requested, symbols_returned, rows_landed, landing_path, landing_sha256)
            VALUES ('y', 'yahoo', ?, '', ?, '', 0, 0, 0, '', '')
            """,
            (D6, D6),
        )
        notes = mock.patch.object(config, "load_yaml", side_effect=lambda f, real=config.load_yaml:
                                  {"notes": []} if f == "break_notes.yaml" else real(f))
        notes.start()
        self.addCleanup(notes.stop)

    def observe(self, as_of, results):
        run_id = f"{NAME}:{as_of}"
        self.conn.execute(
            "INSERT OR IGNORE INTO recon_run (recon_run_id, recon_name, source_a, snapshot_a, source_b, snapshot_b, "
            "window_start, window_end, abs_floor_zar, rel_pct) VALUES (?, ?, 'yahoo', ?, 'eodhd', ?, '', '', 0.01, 0.05)",
            (run_id, NAME, as_of, as_of),
        )
        self.conn.executemany(
            "INSERT INTO recon_result (recon_run_id, key_id, price_date, status, close_a, close_b, diff_zar, diff_pct, "
            "explanation) VALUES (?, ?, ?, ?, 100, 101, 1, 1.0, ?)",
            [(run_id, key, DAY, status, f"{status} on {as_of}") for key, status in results.items()],
        )

    def build(self):
        lifecycle.build(self.conn)
        report.build(self.conn, self.path)
        return load_workbook(self.path)

    def test_tabs_and_rows(self):
        self.observe(D1, {"OPN": "VAL", "TIM": "VAL"})
        self.observe(D2, {"OPN": "VAL", "TIM": "VAL"})
        self.observe(D3, {"OPN": "VAL", "TIM": "MATCH"})
        self.observe(D6, {"OPN": "VAL", "NEW": "UNIT"})
        wb = self.build()
        self.assertEqual(wb.sheetnames, SHEETS)

        ob = wb["Open Breaks"]
        self.assertEqual([ob.cell(r, 2).value for r in range(2, ob.max_row + 1)], ["OPN", "NEW"])
        self.assertEqual((ob["C2"].value, ob["H2"].value), ("Open Ltd", 5))
        self.assertTrue(ob["I2"].value.startswith("=IF(H2<=1"))
        self.assertAlmostEqual(ob["N2"].value, 0.01)  # percentages are stored as fractions
        self.assertEqual(wb["New Today"]["B2"].value, "NEW")
        self.assertEqual([wb["Cleared"].cell(2, c).value for c in (2, 10)], ["TIM", "TIMING"])

    def test_summary_formulas_count_the_columns_they_name(self):
        wb = self.build()
        headers = {ws.title: {c.column_letter: c.value for c in ws[1]} for ws in wb.worksheets[1:]}
        expected = {
            ("Open Breaks", "I"): "Age bucket",
            ("Open Breaks", "J"): "Status",
            ("Open Breaks", "U"): "Cost in a Top 40 fund",
            ("New Today", "B"): "Security",
            ("Cleared", "I"): "Age (trading days)",
            ("Cleared", "J"): "State",
            ("Restatements", "C"): "Status",
            ("Data Quality", "A"): "Check",
            ("Holdings", "B"): "Status",
            ("Holdings", "L"): "Difference (R)",
        }
        seen = set()
        for row in wb["Summary"].iter_rows(min_col=2, max_col=2):
            formula = row[0].value
            if not (isinstance(formula, str) and formula.startswith("=")):
                continue
            for sheet, col in re.findall(r"'?([A-Z][A-Za-z ]+)'?!\$([A-Z])(?:\$\d+)?:", formula):
                self.assertEqual(headers[sheet][col], expected[(sheet, col)], formula)
                seen.add((sheet, col))
        self.assertEqual(seen, set(expected))

    def test_empty_register_still_builds_a_complete_report(self):
        wb = self.build()
        self.assertEqual(wb["Open Breaks"]["A2"].value, f"None as of {D6}")
        self.assertIsNone(wb["Open Breaks"]["B2"].value)  # so the note is never counted as a break
        fonts = {c.font.name for ws in wb for row in ws.iter_rows() for c in row if c.value is not None}
        self.assertEqual(fonts, {"Arial"})


if __name__ == "__main__":
    unittest.main()
