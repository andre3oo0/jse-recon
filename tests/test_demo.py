"""The synthetic demo end to end: every planted error found, every trap quiet, and a complete report written."""

import shutil
import sqlite3
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from src import demo


class DemoTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp())
        cls.sample = cls.tmp / "sample.xlsx"
        cls.lines = demo.build(cls.tmp, cls.sample)
        cls.conn = sqlite3.connect(cls.tmp / "warehouse.db")

    @classmethod
    def tearDownClass(cls):
        cls.conn.close()
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def rows(self, sql):
        return set(self.conn.execute(sql).fetchall())

    def test_the_unit_error_is_the_only_one(self):
        self.assertEqual(self.rows("SELECT security_id, price_date FROM stg_price WHERE unit_anomaly IS NOT NULL"),
                         {demo.PLANTED["unit_error"]})

    def test_moves_the_market_did_not_share_are_flagged_and_the_market_fall_is_not(self):
        moves = self.rows("SELECT source, security_id, price_date FROM v_price_move")
        jump, consolidation, scale = demo.PLANTED["unexplained_jump"], demo.PLANTED["consolidation"], demo.PLANTED["scale_from"]
        self.assertEqual(moves, {(side, *event) for side in demo.SOURCES for event in (jump, consolidation)}
                         | {("synthetic_b", *scale)})

    def test_the_frozen_price_is_found(self):
        code, _ = demo.PLANTED["frozen"]
        self.assertEqual(self.rows("SELECT DISTINCT security_id FROM stg_price WHERE stale_days >= 5"), {(code,)})

    def test_each_planted_break_has_its_type_and_nothing_else_breaks(self):
        opened = self.rows("SELECT key_id, price_date, latest_status FROM break_episode WHERE state = 'OPEN'")
        codes, carried = demo.PLANTED["carried_forward"]
        scale_code, scale_from = demo.PLANTED["scale_from"]
        self.assertTrue({(c, carried, "VAL") for c in codes} <= opened)
        self.assertIn((*demo.PLANTED["unit_error"], "UNIT"), opened)
        self.assertIn((*demo.PLANTED["closed_day"], "CAL"), opened)
        self.assertIn((*demo.PLANTED["restated"], "VAL"), opened)
        self.assertIn((scale_code, scale_from, "SCALE"), opened)
        self.assertEqual({k for k, _, _ in opened}, {"SYB", "SYC", "SYE", "SYF", "SYG", "SYH", "SYI", "SYJ"})

    def test_the_late_day_is_a_timing_difference_that_cleared(self):
        timing = self.rows("SELECT key_id FROM break_episode WHERE state = 'TIMING' AND price_date = '2026-09-10'")
        self.assertEqual(len(timing), 11)  # every share but SYI, whose break carries on as SCALE

    def test_the_restatement_check_finds_the_rewrite_and_the_late_prices(self):
        restated = self.rows(
            "SELECT x.key_id, x.price_date, x.status FROM recon_result x JOIN recon_run r USING (recon_run_id) "
            "WHERE r.recon_name = 'synthetic_a_restatement' AND x.status <> 'MATCH'")
        self.assertIn((*demo.PLANTED["restated"], "VAL"), restated)
        self.assertEqual(sum(1 for _, day, s in restated if s == "ONE_B" and day == demo.LATE_DAY), 12)

    def test_the_consolidation_and_the_small_difference_are_not_errors(self):
        code, day = demo.PLANTED["consolidation"]
        self.assertEqual(self.rows(f"SELECT 1 FROM stg_price WHERE security_id = '{code}' AND unit_anomaly IS NOT NULL"), set())
        small_code, small_day = demo.PLANTED["under_tolerance"]
        self.assertEqual(self.rows(f"SELECT 1 FROM break_episode WHERE key_id = '{small_code}' "
                                   f"AND price_date = '{small_day}'"), set())

    def test_the_answers_say_they_are_synthetic_and_cost_the_late_day(self):
        self.assertTrue(all(line.startswith("SYNTHETIC DATA") for line in self.lines if line.startswith("SYNTHETIC")))
        self.assertEqual(sum(line.startswith("SYNTHETIC DATA") for line in self.lines), 2)
        late = [line for line in self.lines if "was published late" in line]
        self.assertTrue(late and "SYNTHETIC weights" in late[0])

    def test_the_sample_report_is_complete(self):
        wb = load_workbook(self.sample)
        self.assertEqual(wb.sheetnames, ["Summary", "Open Breaks", "New Today", "Cleared", "Restatements",
                                         "Data Quality", "Holdings", "Answers", "About"])
        self.assertGreater(wb["Open Breaks"].max_row, 10)
