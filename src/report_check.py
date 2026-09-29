"""Recalculate a workbook in LibreOffice and fail on any formula error, so the daily report is checked, not trusted."""

import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from src import db, report

ERRORS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A")
MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""


def recalculate(path: Path, timeout: int = 120) -> None:
    # openpyxl writes formulas without results; LibreOffice computes them and saves the values into the file
    with tempfile.TemporaryDirectory(prefix="lo-profile-") as profile:
        url = Path(profile).as_uri()
        subprocess.run(["soffice", "--headless", "--terminate_after_init", f"-env:UserInstallation={url}"],
                       check=True, capture_output=True, timeout=timeout)
        macros = Path(profile) / "user" / "basic" / "Standard"
        if not macros.exists():
            raise RuntimeError("LibreOffice did not create a profile, so nothing was recalculated")
        (macros / "Module1.xba").write_text(MACRO, encoding="utf-8")
        before = path.stat().st_mtime_ns
        subprocess.run(
            ["soffice", "--headless", "--norestore", f"-env:UserInstallation={url}",
             "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application",
             str(path.resolve())],
            check=True, capture_output=True, timeout=timeout,
        )
        if path.stat().st_mtime_ns == before:
            raise RuntimeError("LibreOffice exited without rewriting the file, so nothing was recalculated")


def formula_errors(path: Path) -> tuple[int, list[str]]:
    formulas = load_workbook(path)
    values = load_workbook(path, data_only=True)
    count, errors = 0, []
    for ws in formulas.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    count += 1
                    result = values[ws.title][cell.coordinate].value
                    if isinstance(result, str) and result.startswith(ERRORS):
                        errors.append(f"{ws.title}!{cell.coordinate} = {result}")
    return count, errors


def expected_summary(conn) -> dict[str, int]:
    # Counted straight from the warehouse, independently of the formulas they check
    dq = Counter(r[0] for r in report.dq_rows(conn))
    _, restated = report.restatements(conn)
    status = Counter(r[2] for r in restated)
    states = dict(conn.execute("SELECT state, COUNT(*) FROM break_episode GROUP BY state").fetchall())
    return {
        "Open breaks": states.get("OPEN", 0),
        "New in the latest comparison": len(report.new_breaks(conn)),
        "Cleared as timing differences": states.get("TIMING", 0),
        "Cleared after longer": states.get("CLEARED", 0),
        "Prices rewritten": status["VAL"] + status["UNIT"] + status["STALE"],
        "Prices published late": status["ONE_B"],
        "Prices withdrawn": status["ONE_A"],
        "Unit anomalies": dq["Unit anomaly"],
        "Missing sessions": dq["Missing session"],
        "Missing prices": dq["Missing price"],
        "Frozen prices": dq["Frozen price"],
        "Prices on closed days": dq["Price on closed day"],
    }


def summary_mismatches(path: Path, expected: dict[str, int]) -> list[str]:
    ws = load_workbook(path, data_only=True)["Summary"]
    actual = {r[0]: r[1] for r in ws.iter_rows(min_col=1, max_col=2, values_only=True) if r[0]}
    return [f"{label}: workbook says {actual.get(label)!r}, warehouse says {n}"
            for label, n in expected.items() if actual.get(label) != n]


def main(argv=None) -> int:
    args = list(argv if argv is not None else sys.argv[1:])
    against_warehouse = "--against-warehouse" in args
    paths = [Path(a) for a in args if a != "--against-warehouse"]
    failed = False
    for path in paths:
        recalculate(path)
        count, errors = formula_errors(path)
        print(f"{path.name}: {count} formulas recalculated, {len(errors)} errors")
        for e in errors[:50]:
            print(f"  {e}")
        failed |= bool(errors)
        if against_warehouse:
            mismatches = summary_mismatches(path, expected_summary(db.connect()))
            print(f"{path.name}: Summary checked against the warehouse, {len(mismatches)} mismatches")
            for m in mismatches:
                print(f"  {m}")
            failed |= bool(mismatches)
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
